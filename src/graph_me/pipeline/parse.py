"""Extract plain text from files.

- PDF: pypdfium2. With ``ocr`` (the ``[ocr]`` extra, RapidOCR, models in the wheel, no network),
  pages without a text layer (scans) are read from their image, up to ``OCR_MAX_PAGES``.
- .png / .jpg: only with ``ocr`` and a source's ``images: true`` (photos of documents)
- .enex: Evernote exports, every note's title and text
- .docx / .pptx: read the XML inside the zip with the standard library
- .xlsx: openpyxl (read-only)
- .eml: standard library email parser
- .html: standard library HTML parser, text only
- everything else in TEXT_EXTENSIONS: decoded with charset-normalizer
"""

from __future__ import annotations

import email
import email.policy
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree

from charset_normalizer import from_bytes

MAX_FILE_BYTES = 50 * 1024 * 1024  # skip bigger files (logs, dumps)
MAX_XML_BYTES = 50 * 1024 * 1024  # zip-bomb guard for office files

TEXT_EXTENSIONS = frozenset(
    {
        ".txt", ".md", ".markdown", ".rst", ".org", ".csv", ".tsv", ".json", ".yaml", ".yml",
        ".toml", ".ini", ".cfg", ".log", ".tex", ".xml", ".ics", ".vcf",
        ".py", ".js", ".ts", ".tsx", ".jsx", ".go", ".rs", ".java", ".kt", ".c", ".h", ".cpp",
        ".hpp", ".cs", ".rb", ".php", ".swift", ".sh", ".sql", ".css", ".scss",
    }
)  # fmt: skip
OFFICE_EXTENSIONS = frozenset({".docx", ".pptx", ".xlsx"})
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | OFFICE_EXTENSIONS | {".pdf", ".eml", ".html", ".htm",
                                                              ".enex"}  # fmt: skip
IMAGE_EXTENSIONS = frozenset({".png", ".jpg", ".jpeg"})  # read only with OCR and `images: true`
TABLE_EXTENSIONS = frozenset({".csv", ".tsv"})
OCR_MAX_PAGES = 30  # scanned pages read per PDF: OCR takes about a second a page
OCR_SCALE = 2  # render pages at 144 dpi for OCR

_WESTERN = ("cp1252", "latin_1", "iso8859_15")

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


class ParseError(Exception):
    pass


def is_supported(path: Path, images: bool = False) -> bool:
    suffix = path.suffix.lower()
    return suffix in SUPPORTED_EXTENSIONS or (images and suffix in IMAGE_EXTENSIONS)


def parse_file(path: Path, ocr: bool = False) -> str:
    """Return the text of ``path``. Raises ParseError when the file can't be read."""
    suffix = path.suffix.lower()
    try:
        if path.stat().st_size > MAX_FILE_BYTES:
            raise ParseError(f"larger than {MAX_FILE_BYTES // 1024 // 1024} MB")
        if suffix == ".pdf":
            return _pdf(path, ocr)
        if suffix in IMAGE_EXTENSIONS:
            if not ocr:
                raise ParseError("images are read only with `ocr: true`")
            return _ocr(str(path))
        if suffix == ".enex":
            return _enex(path)
        if suffix == ".docx":
            return _docx(path)
        if suffix == ".pptx":
            return _pptx(path)
        if suffix == ".xlsx":
            return _xlsx(path)
        if suffix == ".eml":
            return parse_email_bytes(path.read_bytes())
        if suffix in {".html", ".htm"}:
            return html_to_text(decode_bytes(path.read_bytes()))
        if suffix in TEXT_EXTENSIONS:
            return decode_bytes(path.read_bytes())
    except ParseError:
        raise
    except Exception as exc:  # corrupt or unexpected files must not stop a scan
        raise ParseError(f"{type(exc).__name__}: {exc}") from exc
    raise ParseError(f"unsupported extension {suffix!r}")


def decode_bytes(data: bytes) -> str:
    if not data:
        return ""
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    matches = list(from_bytes(data))
    if not matches:
        return data.decode("cp1252", errors="replace")
    # Short texts often tie between code pages; prefer Western European ones on a tie.
    top = [
        m
        for m in matches
        if m.chaos == matches[0].chaos and m.coherence == max(x.coherence for x in matches)
    ] or matches[:1]
    for preferred in _WESTERN:
        for m in top:
            if m.encoding == preferred:
                return str(m)
    return str(top[0])


def _pdf(path: Path, ocr: bool = False) -> str:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(path)
    try:
        pages, scanned = [], 0
        for page in pdf:
            textpage = page.get_textpage()
            text = textpage.get_text_bounded()
            textpage.close()
            if not text.strip() and ocr and scanned < OCR_MAX_PAGES:  # a scan: no text layer
                scanned += 1
                bitmap = page.render(scale=OCR_SCALE)
                text = _ocr(bitmap.to_numpy())
                bitmap.close()
            pages.append(text)
            page.close()
        return "\n\n".join(p.strip() for p in pages if p.strip())
    finally:
        pdf.close()


_engine = None


def _ocr(image) -> str:
    """Text lines read in an image (a path or an array) by RapidOCR, one engine per process."""
    global _engine
    if _engine is None:
        try:
            import logging

            from rapidocr import RapidOCR
        except ImportError:
            raise ParseError(
                'scanned pages need the [ocr] extra: uv tool install "graph-me[ocr]"'
            ) from None
        logging.getLogger("RapidOCR").setLevel(logging.WARNING)
        _engine = RapidOCR()
    result = _engine(image)
    return "\n".join(t for t in (getattr(result, "txts", None) or ()) if t and t.strip())


def _enex(path: Path) -> str:
    """Evernote export: one block per note, its title then its text (ENML is XHTML)."""
    root = ElementTree.fromstring(path.read_bytes())
    notes = []
    for note in root.iter("note"):
        title = (note.findtext("title") or "").strip()
        body = html_to_text(note.findtext("content") or "")
        notes.append("\n\n".join(x for x in (title, body.strip()) if x))
    return "\n\n".join(n for n in notes if n)


def _read_member(zf: zipfile.ZipFile, name: str) -> bytes:
    info = zf.getinfo(name)
    if info.file_size > MAX_XML_BYTES:
        raise ParseError(f"{name} too large once uncompressed")
    return zf.read(info)


def _paragraphs(xml: bytes, para_tag: str, text_tag: str) -> list[str]:
    root = ElementTree.fromstring(xml)
    out = []
    for para in root.iter(para_tag):
        text = "".join(t.text or "" for t in para.iter(text_tag)).strip()
        if text:
            out.append(text)
    return out


def _docx(path: Path) -> str:
    with zipfile.ZipFile(path) as zf:
        return "\n\n".join(_paragraphs(_read_member(zf, "word/document.xml"), f"{_W}p", f"{_W}t"))


def _pptx(path: Path) -> str:
    def slide_number(name: str) -> int:
        return int(re.search(r"(\d+)\.xml$", name).group(1))

    with zipfile.ZipFile(path) as zf:
        slides = sorted(
            (n for n in zf.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
            key=slide_number,
        )
        return "\n\n".join(
            "\n".join(_paragraphs(_read_member(zf, s), f"{_A}p", f"{_A}t")) for s in slides
        )


def _xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        parts = []
        for ws in wb.worksheets:
            rows = [
                "\t".join("" if v is None else str(v) for v in row)
                for row in ws.iter_rows(values_only=True)
                if any(v is not None for v in row)
            ]
            if rows:
                parts.append(f"# {ws.title}\n" + "\n".join(rows))
        return "\n\n".join(parts)
    finally:
        wb.close()


def parse_email_bytes(data: bytes) -> str:
    msg = email.message_from_bytes(data, policy=email.policy.default)
    parts = [f"Subject: {msg.get('subject', '')}"]
    for part in msg.walk():
        if part.get_content_maintype() == "multipart" or part.get_filename():
            continue
        if part.get_content_type() == "text/plain":
            parts.append(part.get_content())
        elif part.get_content_type() == "text/html":
            parts.append(html_to_text(part.get_content()))
    return "\n\n".join(p.strip() for p in parts if p and p.strip())


class _TextOnly(HTMLParser):
    _SKIP = {"script", "style", "noscript", "template"}
    _BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self._skipping = 0

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self._skipping += 1
        elif tag in self._BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in self._SKIP and self._skipping:
            self._skipping -= 1

    def handle_data(self, data):
        if not self._skipping:
            self.out.append(data)


def html_to_text(html: str) -> str:
    parser = _TextOnly()
    parser.feed(html)
    parser.close()
    text = "".join(parser.out)
    return re.sub(r"\n\s*\n+", "\n\n", re.sub(r"[ \t]+", " ", text)).strip()
